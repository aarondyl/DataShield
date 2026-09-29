"""DataShield rules assessment, document prefill, and privacy-policy tools."""

import io
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.compliance.policy_checker import check_policy
from app.compliance.policy_generator import generate_policy
from app.compliance.cases_data import cases_for_dimensions
from app.compliance.presets import PRESETS
from app.compliance.questionnaire import QUESTION_MODULES, QUESTIONS, normalize_answers
from app.compliance.report import build_markdown_report, compute_rating, compute_score, level_counts
from app.compliance.roadmap import build_roadmap
from app.compliance.rules import compute_dimension_scores, evaluate_rules
from app.db.session import get_db
from app.models import Assessment, Product
from app.api.developer import upsert_issue
from app.schemas.compliance import (
    AssessmentOut,
    AssessmentRequest,
    DocumentAnalyzeRequest,
    PolicyCheckRequest,
    PolicyGenerateRequest,
)

router = APIRouter(prefix="/compliance", tags=["compliance"])

_SHOW_IF = {
    "data_types": {"key": "collect_personal", "equals": True},
    "sensitive_types": {"key": "collect_sensitive", "equals": True},
    "sensitive_consent": {"key": "collect_sensitive", "equals": True},
    "guardian_consent": {"key": "minors_under_14", "equals": True},
    "cross_border_measure": {"key": "storage_location", "equals": "跨境传输到中国境外"},
    "sdk_disclosed": {"key": "share_third_party", "equals": True},
    "third_party_agreement": {"key": "share_third_party", "equals": True},
    "policy_updated": {"key": "privacy_policy", "equals": True},
}

_HEURISTICS: dict[str, list[str]] = {
    "collect_personal": ["个人信息", "手机号", "邮箱", "注册", "账号", "登录"],
    "collect_sensitive": ["健康", "人脸", "指纹", "生物识别", "定位", "银行卡"],
    "sensitive_consent": ["单独同意"],
    "minors_under_14": ["未成年人", "儿童", "学生"],
    "guardian_consent": ["监护人同意", "家长同意"],
    "encrypt_storage": ["加密", "HTTPS", "TLS", "脱敏"],
    "access_control": ["权限控制", "访问控制", "最小授权", "审计"],
    "breach_plan": ["应急预案", "泄露应急", "安全事件"],
    "retention_defined": ["保存期限", "保留期限", "到期删除", "匿名化"],
    "share_third_party": ["第三方", "SDK", "广告", "统计分析"],
    "sdk_disclosed": ["SDK清单", "第三方清单"],
    "third_party_agreement": ["数据处理协议", "保密协议"],
    "right_access": ["查阅", "复制", "导出"],
    "right_delete": ["注销", "删除账号", "删除个人信息"],
    "right_withdraw": ["撤回同意", "撤回授权"],
    "privacy_policy": ["隐私政策", "隐私协议"],
    "policy_updated": ["更新日期", "最近更新"],
    "consent_popup": ["弹窗", "首次启动"],
    "eu_users": ["欧盟", "欧洲", "出海", "GDPR"],
}


def _question_payload(q: dict[str, Any]) -> dict[str, Any]:
    return {
        "key": q["key"],
        "module": q["module"],
        "type": q["type"],
        "text": q["text"],
        "options": q.get("options", []),
        "default": q.get("default"),
        "show_if": _SHOW_IF.get(q["key"]),
    }


def _assessment_out(item: Assessment) -> AssessmentOut:
    return AssessmentOut(
        id=item.id,
        product_id=item.product_id,
        answers=item.answers,
        hits=item.hits,
        dimension_scores=item.dimension_scores,
        score=item.score,
        rating=item.rating,
        counts=level_counts(item.hits),
        roadmap=build_roadmap(item.hits),
        related_cases=cases_for_dimensions({hit["dimension"] for hit in item.hits})[:6],
        report_markdown=item.report_markdown,
        created_at=item.created_at,
    )


@router.get("/questionnaire")
def questionnaire() -> dict[str, Any]:
    return {
        "modules": QUESTION_MODULES,
        "questions": [_question_payload(q) for q in QUESTIONS],
        "presets": PRESETS,
    }


@router.post("/assessments", response_model=AssessmentOut, status_code=201)
def create_assessment(payload: AssessmentRequest, db: Session = Depends(get_db)) -> AssessmentOut:
    if db.get(Product, payload.product_id) is None:
        raise HTTPException(status_code=404, detail="产品不存在")
    answers = normalize_answers(payload.answers)
    hits = evaluate_rules(answers)
    score = compute_score(hits)
    rating = compute_rating(score)
    dimensions = compute_dimension_scores(hits)
    item = Assessment(
        product_id=payload.product_id,
        answers=answers,
        hits=hits,
        dimension_scores=dimensions,
        score=score,
        rating=rating,
        report_markdown=build_markdown_report(answers, hits, score, rating, dimensions),
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    levels = {"高": "high", "中": "medium", "低": "low"}
    for hit in hits:
        upsert_issue(db, payload.product_id, "quick-check", hit["rule_id"], hit["risk"], levels.get(hit["level"], "medium"), hit["risk"], hit["advice"])
    db.commit()
    return _assessment_out(item)


@router.get("/assessments", response_model=list[AssessmentOut])
def list_assessments(product_id: int | None = None, db: Session = Depends(get_db)) -> list[AssessmentOut]:
    stmt = select(Assessment).order_by(Assessment.id.desc())
    if product_id is not None:
        stmt = stmt.where(Assessment.product_id == product_id)
    return [_assessment_out(item) for item in db.scalars(stmt).all()]


@router.get("/assessments/{assessment_id}", response_model=AssessmentOut)
def get_assessment(assessment_id: int, db: Session = Depends(get_db)) -> AssessmentOut:
    item = db.get(Assessment, assessment_id)
    if item is None:
        raise HTTPException(status_code=404, detail="评估记录不存在")
    return _assessment_out(item)


@router.post("/policy/generate")
def policy_generate(payload: PolicyGenerateRequest) -> dict[str, str]:
    answers = normalize_answers(payload.answers)
    return {"text": generate_policy(answers, payload.product_name, payload.company_name, payload.contact)}


@router.post("/policy/check")
def policy_check(payload: PolicyCheckRequest, db: Session = Depends(get_db)) -> dict[str, Any]:
    result = check_policy(payload.text)
    result["mismatches"] = []
    if payload.product_id is None:
        return result
    product = db.get(Product, payload.product_id)
    if product is None:
        raise HTTPException(status_code=404, detail="产品不存在")
    product.privacy_policy_text = payload.text
    product.has_privacy_policy = True
    comparisons = [
        (product.collects_location_data, "location", ["位置", "定位"], "位置数据未披露", "产品档案显示会收集位置数据，但政策中未发现相关说明。", "补充位置数据的用途、处理方式、权限关闭方式和保存期限。", "位置数据处理"),
        (product.uses_third_party_sdk, "sdk", ["sdk", "第三方服务"], "第三方 SDK 未披露", "产品档案显示使用第三方 SDK，但政策中未发现 SDK 或第三方服务说明。", "逐项披露 SDK 名称、提供方、数据类型、用途和隐私政策链接。", "第三方服务 / SDK 清单"),
        (product.cross_border_data_transfer, "cross-border", ["跨境", "境外", "数据出境"], "跨境传输未披露", "产品档案显示存在跨境传输，但政策中未发现相关说明。", "补充境外接收方、处理目的、数据类型和适用的出境机制。", "信息存储与跨境传输"),
        (product.children_related, "children", ["儿童", "未成年人", "监护人"], "未成年人处理未披露", "产品档案显示涉及未成年人，但政策中未发现专项说明。", "补充年龄范围、监护人同意和未成年人权利保障方式。", "未成年人保护"),
    ]
    lower = payload.text.lower()
    for enabled, key, words, title, why, fix, placement in comparisons:
        if enabled and not any(word.lower() in lower for word in words):
            mismatch = {"key": key, "title": title, "product_behavior": why.split("，但")[0], "policy": "未发现相关说明", "advice": fix}
            result["mismatches"].append(mismatch)
            upsert_issue(db, product.id, "policy-check", key, title, "high" if key in {"location","children"} else "medium", why, fix, "请根据产品实际处理活动补充完整、准确的披露。", placement)
    for item in result["items"]:
        if not item["covered"]:
            key = "element-" + item["name"]
            upsert_issue(db, product.id, "policy-check", key, f'补充隐私政策要素：{item["name"]}', "medium", f'当前隐私政策未覆盖“{item["name"]}”，用户无法充分了解相关处理规则。', item["hint"], '', item["name"])
    db.commit()
    return result


def _analyze_text(text: str) -> dict[str, Any]:
    suggestions: dict[str, Any] = {}
    evidence: dict[str, Any] = {}
    for key, keywords in _HEURISTICS.items():
        for keyword in keywords:
            pos = text.lower().find(keyword.lower())
            if pos >= 0:
                suggestions[key] = True
                start, end = max(0, pos - 30), min(len(text), pos + len(keyword) + 30)
                evidence[key] = {"keyword": keyword, "snippet": text[start:end].replace("\n", " ")}
                break
    if any(word.lower() in text.lower() for word in ["境外服务器", "跨境", "AWS", "Azure"]):
        suggestions["storage_location"] = "跨境传输到中国境外"
    return {
        "suggestions": suggestions,
        "evidence": evidence,
        "mode": "heuristic",
        "note": "已根据文档内容生成问卷预填建议，请逐项确认。",
    }


@router.post("/documents/analyze")
def analyze_document(payload: DocumentAnalyzeRequest) -> dict[str, Any]:
    return _analyze_text(payload.text)


@router.post("/documents/upload")
async def upload_document(file: UploadFile = File(...)) -> dict[str, Any]:
    suffix = Path(file.filename or "").suffix.lower()
    data = await file.read()
    try:
        if suffix in {".txt", ".md"}:
            text = data.decode("utf-8", errors="ignore")
        elif suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif suffix == ".docx":
            import docx

            document = docx.Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            parts.extend(" ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            text = "\n".join(parts)
        else:
            raise HTTPException(status_code=400, detail="仅支持 .txt、.md、.pdf、.docx 文件")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"文档解析失败：{exc}") from exc
    if not text.strip():
        raise HTTPException(status_code=400, detail="文档中没有可提取的文本")
    return {**_analyze_text(text), "filename": file.filename, "characters": len(text)}
