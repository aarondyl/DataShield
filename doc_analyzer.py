# -*- coding: utf-8 -*-
"""文档智能分析模块：上传项目计划书，自动预填合规问卷。

支持 .txt / .md / .docx / .pdf 四种格式，也可以直接粘贴文本。
分析提供两种模式：
    AI 模式：交给 LLM 通读文档后代为作答（配置 API Key 时可用）；
    规则模式：关键词启发式打分（离线可用），同时给出每个判断的匹配依据，
             方便用户核对预填结果是否可靠。
两种模式都只给出「预填建议」，最终答案由用户在问卷页逐项确认。
"""

import io

from llm_explainer import analyze_document, is_llm_available
from questionnaire import QUESTIONS

# 传给 LLM 的文档最大长度（字符），超长截断以控制 token 消耗
MAX_TEXT_FOR_LLM = 6000

# 关键词启发式：题目 key -> (关键词组, 命中时建议值)
# 词组内任意关键词命中即视为「文档提到了该事项」
HEURISTIC_RULES = {
    "collect_personal": (["个人信息", "手机号", "邮箱", "注册", "账号", "登录", "用户名"], True),
    "collect_sensitive": (["健康", "人脸", "指纹", "生物识别", "定位", "位置", "银行卡", "支付账户"], True),
    "sensitive_consent": (["单独同意"], True),
    "minors_under_14": (["未成年人", "儿童", "青少年", "学生"], True),
    "guardian_consent": (["监护人同意", "家长同意"], True),
    "encrypt_storage": (["加密", "https", "HTTPS", "TLS", "脱敏"], True),
    "access_control": (["权限控制", "访问控制", "最小授权", "审计"], True),
    "breach_plan": (["应急预案", "泄露应急", "安全事件"], True),
    "retention_defined": (["保存期限", "保留期限", "到期删除", "匿名化"], True),
    "share_third_party": (["第三方", "SDK", "sdk", "广告", "统计分析", "友盟", "极光"], True),
    "sdk_disclosed": (["SDK清单", "sdk清单", "第三方清单"], True),
    "third_party_agreement": (["数据处理协议", "保密协议"], True),
    "right_access": (["查阅", "复制", "导出"], True),
    "right_delete": (["注销", "删除账号", "删除个人信息"], True),
    "right_withdraw": (["撤回同意", "撤回授权"], True),
    "privacy_policy": (["隐私政策", "隐私协议", "用户隐私"], True),
    "policy_updated": (["更新日期", "最近更新"], True),
    "consent_popup": (["弹窗", "首次启动", "启动页提示"], True),
    "eu_users": (["欧盟", "欧洲", "海外用户", "出海", "GDPR"], True),
}

# 存储位置推断词（单独处理，因为答案是选项而非布尔）
STORAGE_HINTS = {
    "跨境传输到中国境外": ["境外服务器", "海外服务器", "跨境", "AWS", "亚马逊云", "谷歌云", "Azure"],
    "欧盟境内": ["欧盟服务器", "欧洲服务器", "法兰克福节点"],
}


def extract_text(uploaded_file):
    """从上传的文件提取纯文本。

    参数：uploaded_file: Streamlit UploadedFile 对象。
    返回：str 文本；解析失败抛出异常（由页面捕获并提示）。
    """
    name = uploaded_file.name.lower()
    data = uploaded_file.read()

    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="ignore")

    if name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)

    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    raise ValueError("不支持的文件格式，请上传 .txt / .md / .docx / .pdf")


def questions_brief():
    """生成给 LLM 的题目简述（不含 show_if，全量列出）。"""
    return [
        {"key": q["key"], "text": q["text"], "type": q["type"], "options": q.get("options", [])}
        for q in QUESTIONS
    ]


def _find_snippet(text, keyword, width=30):
    """截取关键词在文中的上下文片段，作为「匹配依据」展示。"""
    idx = text.find(keyword)
    if idx < 0:
        return ""
    start = max(0, idx - width)
    end = min(len(text), idx + len(keyword) + width)
    snippet = text[start:end].replace("\n", " ")
    return ("…" if start > 0 else "") + snippet + ("…" if end < len(text) else "")


def heuristic_analyze(text):
    """关键词启发式分析（离线模式）。

    返回：(suggestions, evidence)
        suggestions: {题目 key: 建议答案}
        evidence:    {题目 key: {"keyword": 命中的关键词, "snippet": 上下文片段}}
    """
    suggestions, evidence = {}, {}
    for key, (keywords, value) in HEURISTIC_RULES.items():
        for kw in keywords:
            if kw in text:
                suggestions[key] = value
                evidence[key] = {"keyword": kw, "snippet": _find_snippet(text, kw)}
                break
    for location, keywords in STORAGE_HINTS.items():
        for kw in keywords:
            if kw in text:
                suggestions["storage_location"] = location
                evidence["storage_location"] = {"keyword": kw, "snippet": _find_snippet(text, kw)}
                break
        if "storage_location" in suggestions:
            break
    return suggestions, evidence


def analyze(text):
    """文档分析主入口：优先 AI 模式，失败/未配置时自动用启发式。

    返回：dict：
        {
          "suggestions": {题目 key: 建议答案},
          "evidence":    {题目 key: 匹配依据}（仅启发式模式有），
          "mode": "llm" 或 "heuristic",
          "truncated": 文档是否因超长被截断,
          "note": 给用户的说明文字,
        }
    """
    truncated = len(text) > MAX_TEXT_FOR_LLM
    if is_llm_available():
        result = analyze_document(text[:MAX_TEXT_FOR_LLM], questions_brief())
        if result:
            # 过滤掉题目之外的 key，避免脏数据进入问卷
            valid_keys = {q["key"] for q in QUESTIONS}
            suggestions = {k: v for k, v in result.items() if k in valid_keys}
            if suggestions:
                note = "AI 已通读文档并给出预填建议，请在问卷页逐项确认后再提交。"
                if truncated:
                    note += f"文档较长，仅分析了前 {MAX_TEXT_FOR_LLM} 字。"
                return {
                    "suggestions": suggestions,
                    "evidence": {},
                    "mode": "llm",
                    "truncated": truncated,
                    "note": note,
                }
    suggestions, evidence = heuristic_analyze(text)
    return {
        "suggestions": suggestions,
        "evidence": evidence,
        "mode": "heuristic",
        "truncated": False,
        "note": "已用关键词规则给出预填建议（未配置 API Key 或 AI 调用失败时的离线模式），"
                "覆盖不全属正常，请在问卷页逐项确认。",
    }
